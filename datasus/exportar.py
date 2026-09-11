# -*- coding: utf-8 -*-
"""Exportação da tabela de estabelecimentos para Excel e PDF."""
from __future__ import annotations

import io

import pandas as pd


def adicionar_linha_total(df: pd.DataFrame, coluna_total: str) -> pd.DataFrame:
    """Devolve uma cópia do DataFrame com uma linha "Total" ao final,
    somando `coluna_total` e deixando as demais colunas em branco."""
    if df.empty:
        return df
    total = {col: "" for col in df.columns}
    total[df.columns[0]] = "Total"
    total[coluna_total] = df[coluna_total].sum()
    return pd.concat([df, pd.DataFrame([total])], ignore_index=True)


def _linha_filtros(filtros: dict[str, str] | None) -> str:
    """Formata o dicionário de filtros aplicados como uma linha de texto
    legível (ex.: 'Tipo: Transplante · Região: Nordeste, Sul · UF: CE')."""
    if not filtros:
        return "Nenhum filtro aplicado (todos os dados)."
    partes = [f"{chave}: {valor}" for chave, valor in filtros.items() if valor]
    return " · ".join(partes) if partes else "Nenhum filtro aplicado (todos os dados)."


def gerar_excel(
    df: pd.DataFrame,
    titulo: str = "Estabelecimentos",
    filtros: dict[str, str] | None = None,
) -> bytes:
    """Gera um .xlsx formatado (cabeçalho em negrito, colunas
    autoajustadas), com os filtros aplicados listados nas primeiras
    linhas da planilha, e devolve os bytes prontos para download."""
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    buffer = io.BytesIO()
    aba = titulo[:31] or "Dados"
    linha_filtros = _linha_filtros(filtros)

    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        # Duas linhas de cabeçalho (título + filtros aplicados) antes da
        # tabela propriamente dita, para o arquivo exportado ser
        # autoexplicativo sobre qual recorte de dados ele representa.
        df.to_excel(writer, index=False, sheet_name=aba, startrow=3)
        planilha = writer.sheets[aba]
        planilha.cell(row=1, column=1, value=titulo).font = Font(bold=True, size=13)
        planilha.cell(row=2, column=1, value=f"Filtros aplicados: {linha_filtros}").font = Font(
            italic=True, size=9, color="6E6E73"
        )

        cabecalho_linha = 4
        for col_idx, coluna in enumerate(df.columns, start=1):
            celula = planilha.cell(row=cabecalho_linha, column=col_idx)
            celula.font = Font(bold=True)
            largura = max(12, min(45, int(df[coluna].astype(str).str.len().max() or 12) + 2))
            planilha.column_dimensions[get_column_letter(col_idx)].width = largura
    return buffer.getvalue()


def gerar_pdf(
    df: pd.DataFrame,
    titulo: str = "Estabelecimentos",
    filtros: dict[str, str] | None = None,
) -> bytes:
    """Gera um PDF simples em tabela (retrato/paisagem automático
    conforme o número de colunas), com os filtros aplicados listados no
    topo do documento, e devolve os bytes."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet

    buffer = io.BytesIO()
    pagesize = landscape(A4) if len(df.columns) > 5 else A4
    doc = SimpleDocTemplate(
        buffer, pagesize=pagesize,
        leftMargin=1.5 * cm, rightMargin=1.5 * cm,
        topMargin=1.5 * cm, bottomMargin=1.5 * cm,
    )
    estilos = getSampleStyleSheet()
    estilo_filtros = estilos["Normal"].clone("filtros")
    estilo_filtros.textColor = colors.HexColor("#6E6E73")
    estilo_filtros.fontSize = 9

    linha_filtros = _linha_filtros(filtros)
    elementos = [
        Paragraph(titulo, estilos["Title"]),
        Paragraph(f"Filtros aplicados: {linha_filtros}", estilo_filtros),
        Spacer(1, 12),
    ]

    dados_tabela = [list(df.columns)] + df.astype(str).values.tolist()
    tabela = Table(dados_tabela, repeatRows=1)
    tabela.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0B3D66")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D2D2D7")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F5F7")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    elementos.append(tabela)
    doc.build(elementos)
    return buffer.getvalue()
